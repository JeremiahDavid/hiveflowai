"""Read ``.xlsx`` with openpyxl, preserving merges, number formats, and styles."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from hiveflow_spreadsheet_parser.readers.base import (
    Cell,
    CellKind,
    SheetGrid,
    Workbook,
    col_index,
)

# openpyxl ``data_type`` -> our CellKind (with ``data_only=True`` cached values)
_KIND: dict[str, CellKind] = {
    "n": "number",
    "s": "text",
    "str": "text",
    "d": "date",
    "b": "bool",
    "f": "formula",
    "e": "error",
}


def _style_key(cell: Any) -> int:
    font = cell.font
    fill = cell.fill
    fg = getattr(fill, "fgColor", None)
    filled = bool(
        fill
        and fill.patternType
        and fg is not None
        and getattr(fg, "rgb", None) not in (None, "00000000")
    )
    align = cell.alignment.horizontal or ""
    return (
        hash(
            (
                bool(font.bold),
                bool(font.italic),
                round(float(font.size or 11)),
                filled,
                align,
            )
        )
        & 0x7FFFFFFF
    )


def _classify(cell: Any) -> tuple[CellKind, object]:
    value = cell.value
    if value is None:
        return "empty", None
    if isinstance(value, str):
        if value.strip() == "":
            return "empty", None
        return "text", value
    kind = _KIND.get(cell.data_type, "text")
    if kind == "formula":  # data_only workbook with no cached value
        return "empty", None
    return kind, value


def _frozen(ws: Worksheet) -> tuple[int, int]:
    fp = ws.freeze_panes
    if not fp:
        return 0, 0
    from hiveflow_spreadsheet_parser.readers.base import parse_a1_cell

    r0, c0 = parse_a1_cell(fp)
    return r0, c0


def _sheet_grid(ws: Worksheet) -> SheetGrid:
    max_row = ws.max_row or 1
    max_col = ws.max_column or 1

    merged: list[tuple[int, int, int, int]] = []
    merged_members: set[tuple[int, int]] = set()
    merge_origins: set[tuple[int, int]] = set()
    for rng in ws.merged_cells.ranges:
        r0, c0, r1, c1 = rng.min_row - 1, rng.min_col - 1, rng.max_row - 1, rng.max_col - 1
        merged.append((r0, c0, r1, c1))
        merge_origins.add((r0, c0))
        for r in range(r0, r1 + 1):
            for c in range(c0, c1 + 1):
                merged_members.add((r, c))

    cells: list[list[Cell]] = []
    last_row = -1
    last_col = -1
    for r_idx, row in enumerate(
        ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col)
    ):
        out_row: list[Cell] = []
        for c_idx, oc in enumerate(row):
            kind, value = _classify(oc)
            cell = Cell(
                value=value,
                kind=kind,
                number_format=(oc.number_format if kind != "empty" else None),
                is_merged=(r_idx, c_idx) in merged_members,
                merge_origin=(r_idx, c_idx) in merge_origins,
                style_key=_style_key(oc) if kind != "empty" else 0,
            )
            out_row.append(cell)
            if kind != "empty":
                last_row = max(last_row, r_idx)
                last_col = max(last_col, c_idx)
        cells.append(out_row)

    # Trim trailing all-empty rows/cols (openpyxl often over-reports dimensions).
    nrows = max(last_row + 1, 1)
    ncols = max(last_col + 1, 1)
    cells = [row[:ncols] + [Cell()] * max(0, ncols - len(row)) for row in cells[:nrows]]

    hidden_rows = {i - 1 for i, dim in ws.row_dimensions.items() if dim.hidden and i - 1 < nrows}
    hidden_cols = {
        col_index(letter)
        for letter, dim in ws.column_dimensions.items()
        if dim.hidden and col_index(letter) < ncols
    }
    frozen_rows, frozen_cols = _frozen(ws)

    return SheetGrid(
        name=ws.title,
        nrows=nrows,
        ncols=ncols,
        cells=cells,
        merged_ranges=[m for m in merged if m[0] < nrows and m[1] < ncols],
        hidden_rows=hidden_rows,
        hidden_cols=hidden_cols,
        frozen_rows=frozen_rows,
        frozen_cols=frozen_cols,
    )


class XlsxReader:
    """`WorkbookReader` for ``.xlsx`` / ``.xlsm`` files."""

    extensions = (".xlsx", ".xlsm")

    def read(self, path: str | Path) -> Workbook:
        p = Path(path)
        wb = load_workbook(p, data_only=True, read_only=False)
        try:
            sheets = [_sheet_grid(wb[name]) for name in wb.sheetnames]
        finally:
            wb.close()
        return Workbook(path=str(p), fmt="xlsx", sheets=sheets)


def read_workbook(path: str | Path) -> Workbook:
    """Read ``path`` with the reader matching its extension."""
    p = Path(path)
    ext = p.suffix.lower()
    for reader in (XlsxReader(),):
        if ext in reader.extensions:
            return reader.read(p)
    raise ValueError(f"no reader for {ext!r}; supported: .xlsx .xlsm")
