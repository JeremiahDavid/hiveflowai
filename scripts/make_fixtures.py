"""Generate messy .xlsx fixtures for the spreadsheet-parser test suite.

    python scripts/make_fixtures.py

Each workbook targets one real-world layout problem the parser must survive.
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

FIXTURES = Path(__file__).resolve().parents[1] / ("packages/spreadsheet-parser/tests/fixtures")

BOLD = Font(bold=True)


def _header(ws: Worksheet, row: int, start_col: int, names: list[str]) -> None:
    for i, name in enumerate(names):
        cell = ws.cell(row=row, column=start_col + i, value=name)
        cell.font = BOLD


def clean_single(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Orders"
    _header(ws, 1, 1, ["order_id", "customer", "region", "qty", "amount"])
    for i in range(20):
        ws.append([1000 + i, f"Cust {i % 5}", ["N", "S", "E", "W"][i % 4], i + 1, (i + 1) * 12.5])
    wb.save(path)


def two_stacked(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    _header(ws, 1, 1, ["sku", "name", "on_hand"])
    for i in range(4):
        ws.append([f"SKU{i}", f"Item {i}", i * 10])
    ws["A7"] = None
    _header(ws, 10, 1, ["month", "revenue", "cost", "margin"])
    for i in range(6):
        r = 11 + i
        ws.cell(row=r, column=1, value=f"2026-{i + 1:02d}")
        ws.cell(row=r, column=2, value=1000 + i * 100)
        ws.cell(row=r, column=3, value=600 + i * 50)
        ws.cell(row=r, column=4, value=400 + i * 50)
    wb.save(path)


def report_export(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales Report"
    ws["A1"] = "ACME CORP — Confidential"
    ws["A2"] = "Sales by Rep"
    ws["A3"] = "Period: 2026-Q1"
    ws["A4"] = "Generated 2026-04-01 08:12"
    _header(ws, 5, 1, ["rep", "deals", "bookings_usd"])
    reps = ["Alice", "Bob", "Carol", "Dan", "Eve"]
    total = 0.0
    for i, rep in enumerate(reps):
        bookings = 12000 + i * 3500
        total += bookings
        ws.cell(row=6 + i, column=1, value=rep)
        ws.cell(row=6 + i, column=2, value=3 + i)
        ws.cell(row=6 + i, column=3, value=bookings)
    ws.cell(row=12, column=1, value="Grand Total")
    ws.cell(row=12, column=2, value=sum(3 + i for i in range(5)))
    ws.cell(row=12, column=3, value=total)
    wb.save(path)


def alternating_blank(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Export"
    _header(ws, 1, 1, ["ticket", "status", "hours"])
    row = 2
    for i in range(12):
        ws.cell(row=row, column=1, value=f"T-{200 + i}")
        ws.cell(row=row, column=2, value=["open", "closed", "pending"][i % 3])
        ws.cell(row=row, column=3, value=round((i + 1) * 1.25, 2))
        row += 2  # blank spacer row between every record
    wb.save(path)


def wrapped_records(path: Path) -> None:
    """Each record's fields are split across a *variable* number of physical rows.

    Row 1 is the header. Every record starts with a row where ``id`` is populated;
    its remaining fields trail across however many rows follow, one continuation
    row of a fixed size would not cover this — record lengths here are 1, 2, 3, 4.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Timesheet"
    _header(ws, 1, 1, ["id", "name", "dept", "hours"])
    r = 2
    ws.cell(row=r, column=1, value="W-1")
    r += 1
    ws.cell(row=r, column=2, value="Alice")
    ws.cell(row=r, column=3, value="Eng")
    ws.cell(row=r, column=4, value=5)
    r += 1
    ws.cell(row=r, column=1, value="W-2")
    r += 1
    ws.cell(row=r, column=2, value="Bob")
    r += 1
    ws.cell(row=r, column=3, value="Sales")
    ws.cell(row=r, column=4, value=3)
    r += 1
    ws.cell(row=r, column=1, value="W-3")
    ws.cell(row=r, column=2, value="Carol")
    ws.cell(row=r, column=3, value="Ops")
    ws.cell(row=r, column=4, value=8)
    r += 1
    ws.cell(row=r, column=1, value="W-4")
    r += 1
    ws.cell(row=r, column=2, value="Dave")
    r += 1
    ws.cell(row=r, column=3, value="Support")
    r += 1
    ws.cell(row=r, column=4, value=2)
    wb.save(path)


def merged_headers(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.cell(row=1, column=2, value="Q1").font = BOLD
    ws.cell(row=1, column=4, value="Q2").font = BOLD
    ws.merge_cells("B1:C1")
    ws.merge_cells("D1:E1")
    _header(ws, 2, 1, ["category", "plan", "actual", "plan", "actual"])
    groups = [("Marketing", 3), ("Sales", 3)]
    r = 3
    for name, n in groups:
        ws.cell(row=r, column=1, value=name)
        ws.merge_cells(start_row=r, start_column=1, end_row=r + n - 1, end_column=1)
        for k in range(n):
            ws.cell(row=r + k, column=2, value=100 + k)
            ws.cell(row=r + k, column=3, value=90 + k)
            ws.cell(row=r + k, column=4, value=110 + k)
            ws.cell(row=r + k, column=5, value=105 + k)
        r += n
    wb.save(path)


def pivot_matrix(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Pivot"
    regions = ["North", "South", "East", "West"]
    months = ["Jan", "Feb", "Mar"]
    ws.cell(row=1, column=1, value="Region")
    for j, m in enumerate(months):
        ws.cell(row=1, column=2 + j, value=m).font = BOLD
    ws.cell(row=1, column=2 + len(months), value="Grand Total").font = BOLD
    for i, reg in enumerate(regions):
        r = 2 + i
        ws.cell(row=r, column=1, value=reg)
        rowtot = 0
        for j in range(len(months)):
            v = (i + 1) * 100 + j * 10
            rowtot += v
            ws.cell(row=r, column=2 + j, value=v)
        ws.cell(row=r, column=2 + len(months), value=rowtot)
    gr = 2 + len(regions)
    ws.cell(row=gr, column=1, value="Grand Total")
    for j in range(len(months) + 1):
        ws.cell(
            row=gr,
            column=2 + j,
            value=sum(ws.cell(row=2 + i, column=2 + j).value or 0 for i in range(len(regions))),
        )
    wb.save(path)


def side_by_side(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "TwoUp"
    _header(ws, 1, 1, ["dept", "headcount", "budget"])
    _header(ws, 1, 6, ["vendor", "spend", "contracts"])
    for i in range(10):
        ws.cell(row=2 + i, column=1, value=f"Dept {i}")
        ws.cell(row=2 + i, column=2, value=5 + i)
        ws.cell(row=2 + i, column=3, value=50000 + i * 1000)
        ws.cell(row=2 + i, column=6, value=f"Vendor {i}")
        ws.cell(row=2 + i, column=7, value=1000 * (i + 1))
        ws.cell(row=2 + i, column=8, value=i % 3 + 1)
    wb.save(path)


def formats(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoices"
    _header(ws, 1, 1, ["invoice_id", "  client  ", "issued", "amount", "tax_rate", "paid"])
    base = dt.date(2026, 1, 5)
    rows = [
        (5001, "  Globex ", base, 1240.00, 0.08, True),
        (5002, "Initech", base + dt.timedelta(days=9), 980.50, 0.08, False),
        (5003, " Umbrella  ", base + dt.timedelta(days=15), 15230.00, 0.00, True),
        (5004, "Hooli", base + dt.timedelta(days=22), 4300.25, 0.075, False),
        (5003, " Umbrella  ", base + dt.timedelta(days=15), 15230.00, 0.00, True),  # dup
    ]
    r = 2
    for inv, client, issued, amount, tax, paid in rows:
        ws.cell(row=r, column=1, value=inv)
        ws.cell(row=r, column=2, value=client)
        c = ws.cell(row=r, column=3, value=issued)
        c.number_format = "yyyy-mm-dd"
        c = ws.cell(row=r, column=4, value=amount)
        c.number_format = '"$"#,##0.00'
        c = ws.cell(row=r, column=5, value=tax)
        c.number_format = "0.0%"
        ws.cell(row=r, column=6, value=paid)
        r += 1
    r += 1  # fully blank row
    ws.cell(row=r, column=1, value="TOTAL")
    ws.cell(row=r, column=4, value=sum(x[3] for x in rows[:-1]))
    wb.save(path)


BUILDERS = {
    "clean_single.xlsx": clean_single,
    "two_stacked.xlsx": two_stacked,
    "report_export.xlsx": report_export,
    "alternating_blank.xlsx": alternating_blank,
    "wrapped_records.xlsx": wrapped_records,
    "merged_headers.xlsx": merged_headers,
    "pivot_matrix.xlsx": pivot_matrix,
    "side_by_side.xlsx": side_by_side,
    "formats.xlsx": formats,
}


def main() -> None:
    FIXTURES.mkdir(parents=True, exist_ok=True)
    for name, builder in BUILDERS.items():
        builder(FIXTURES / name)
        print(f"wrote {name}")


if __name__ == "__main__":
    main()
